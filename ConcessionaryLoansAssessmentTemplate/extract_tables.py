"""
Module for extracting and transforming Excel tables from XLSX workbooks.

This module uses openpyxl and pandas to read Excel Tables from a workbook,
apply configurable transformation functions, and output the final transformed
data as CSV files.
"""

import csv
import os
import re

import openpyxl
import pandas as pd

INPUT_DIR = "input"
OUTPUT_DIR = "output"
FILENAME = "ConcessionaryLoansApplicationTemplatev2_8.xlsx"
WORKBOOK_PATH = os.path.join(INPUT_DIR, FILENAME)
YEAR_PATTERN = re.compile(r"^2\d{3}$")  # Matches strings like '2025', '2070', etc.


def unpivot_table(
    df, identifier_columns, key_column_name, value_column_name, table_name
):
    """
    Unpivots a DataFrame using pandas.melt.

    Parameters:
      identifier_columns (list): Columns to keep as identifiers (e.g., line item,
          value basis, base year, comment).
      key_column_name (str): Name for the new column that will hold the original
          column headers (e.g., Year).
      value_column_name (str): Name for the new column that will hold the
          corresponding values (e.g., Cost).

    Returns:
      pd.DataFrame: The unpivoted (melted) DataFrame.
    """
    identifier_cols_existing = [col for col in identifier_columns if col in df.columns]
    if set(identifier_cols_existing) != set(identifier_columns):
        missing = set(identifier_columns) - set(identifier_cols_existing)
        raise ValueError(f"Columns {missing} not found in DataFrame.")
    value_columns = [col for col in df.columns if YEAR_PATTERN.match(col)]
    extra_columns = [
        col for col in df.columns if col not in identifier_cols_existing + value_columns
    ]
    if extra_columns:
        raise ValueError(
            f"Unexpected columns found in table {table_name}: {extra_columns}"
        )
    df_melted = pd.melt(
        df,
        id_vars=identifier_cols_existing,
        value_vars=value_columns,
        var_name=key_column_name,
        value_name=value_column_name,
    )
    df_melted = df_melted.dropna(subset=[value_column_name])
    return df_melted


# Master dictionary mapping table names to sequences of transformation steps.
# Each step is defined as a dictionary with an 'action' (transformation function)
# and a 'kwargs' dictionary containing its arguments.
table_transformations = {
    "Opex_NonFuel_PP": [
        {
            "action": unpivot_table,
            "kwargs": {
                "identifier_columns": [
                    "Opex item",
                    "Value basis",
                    "Base year",
                    "Applicant comment",
                    "ProjectId",
                    "LifeCycleMonths",
                    "Scenario",
                    "ExpenseType",
                ],
                "key_column_name": "Year",
                "value_column_name": "Cost",
            },
        }
    ],
    "Opex_NonFuel_BC": [
        {
            "action": unpivot_table,
            "kwargs": {
                "identifier_columns": [
                    "ProjectId",
                    "LifeCycleMonths",
                    "Opex item",
                    "Scenario",
                    "Value basis",
                    "Base year",
                    "Applicant comment",
                    "ExpenseType",
                ],
                "key_column_name": "Year",
                "value_column_name": "Cost",
            },
        }
    ],
    "Capex_PP": [
        {
            "action": unpivot_table,
            "kwargs": {
                "identifier_columns": [
                    "ProjectId",
                    "LifeCycleMonths",
                    "Capex item",
                    "Scenario",
                    "Value basis",
                    "Base year",
                    "Applicant comment",
                    "ExpenseType",
                ],
                "key_column_name": "Year",
                "value_column_name": "Cost",
            },
        }
    ],
    "Capex_BC": [
        {
            "action": unpivot_table,
            "kwargs": {
                "identifier_columns": [
                    "ProjectId",
                    "LifeCycleMonths",
                    "Capex item",
                    "Scenario",
                    "Value basis",
                    "Base year",
                    "Applicant comment",
                    "ExpenseType",
                ],
                "key_column_name": "Year",
                "value_column_name": "Cost",
            },
        }
    ],
    "FuelUse_PP": [
        {
            "action": unpivot_table,
            "kwargs": {
                "identifier_columns": [
                    "ProjectId",
                    "LifeCycleMonths",
                    "Scenario",
                    "Fuel type",
                    "Units",
                    "Applicant comment",
                ],
                "key_column_name": "Year",
                "value_column_name": "Cost",
            },
        }
    ],
    "FuelUse_BC": [
        {
            "action": unpivot_table,
            "kwargs": {
                "identifier_columns": [
                    "ProjectId",
                    "LifeCycleMonths",
                    "Scenario",
                    "Fuel type",
                    "Units",
                    "Applicant comment",
                ],
                "key_column_name": "Year",
                "value_column_name": "Cost",
            },
        }
    ],
}


def extract_table_data(sheet, table):
    """
    Extracts data from a given Excel table within a sheet.

    Parameters:
      sheet: The worksheet object.
      table: The table object with a 'ref' attribute.

    Returns:
      tuple: A tuple (header, rows) where header is the first row of the table
             and rows is a list of the remaining rows.
    """
    table_range = table.ref
    data = []
    for row in sheet[table_range]:
        data.append([cell.value for cell in row])
    header = data[0]
    rows = data[1:]
    return header, rows


def extract_tables_from_workbook(xlsx_path):
    """
    Opens the provided XLSX workbook, iterates through each worksheet,
    finds Excel Tables, and for each table:
      - If a transformation is configured, converts the raw data to a DataFrame,
        applies the transformation functions in sequence, and writes out the
        final CSV.
      - Otherwise, writes out the raw table data as CSV.
    """
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    for sheet in wb.worksheets:
        if hasattr(sheet, "tables") and sheet.tables:
            for table in sheet.tables.values():
                table_range = table.ref
                print(
                    f"Extracting table '{table.name}' ",
                    f"from sheet '{sheet.title}' ",
                    f"with range {table_range}",
                )
                header, rows = extract_table_data(sheet, table)
                df = pd.DataFrame(rows, columns=header)
                if table.name in table_transformations:
                    for transform in table_transformations[table.name]:
                        action_func = transform["action"]
                        kwargs = transform.get("kwargs", {})
                        kwargs["table_name"] = table.name
                        df = action_func(df, **kwargs)
                    csv_filename = f"{table.name}.csv"
                    csv_path = os.path.join(OUTPUT_DIR, csv_filename)
                    df.to_csv(csv_path, index=False)
                    print(
                        f"Transformed table '{table.name}' saved to '{csv_filename}'."
                    )
                else:
                    csv_filename = f"{table.name}.csv"
                    csv_path = os.path.join(OUTPUT_DIR, csv_filename)
                    with open(csv_path, "w", newline="", encoding="utf-8") as f:
                        writer = csv.writer(f)
                        writer.writerow(header)
                        writer.writerows(rows)
                    print(f"Table '{table.name}' extracted to '{csv_filename}'.")
    print("Extraction complete.")


if __name__ == "__main__":
    extract_tables_from_workbook(WORKBOOK_PATH)
